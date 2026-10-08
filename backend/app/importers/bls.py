"""Importer für den Bundeslebensmittelschlüssel (BLS 4.0, Max Rubner-Institut).

Liest die lokal heruntergeladene Datendatei (`BLS_4_0_Daten_2025_DE.xlsx`, direkt oder in der
ZIP-Datei `BLS_4_0_2025_DE.zip`) und liefert `IngredientRecord`s mit Nährwerten je 100 g.
Reine Logik ohne Datenbankbezug. Die BLS-Dateien liegen nie im Repo.

Dateiformat: Zeile 1 ist die Kopfzeile. Spalte A `BLS Code` (Buchstabe + 6 Ziffern, der Buchstabe
ist die Warengruppe), B `Lebensmittelbezeichnung`, C `Food name`. Ab Spalte D gibt es je Nährstoff
drei Spalten: Wert (Kopf `CODE Bezeichnung [Einheit/100g]`), `CODE Datenherkunft`, `CODE Referenz`.
Der Nährstoffcode ist das erste Wort der Kopfzeile, die Einheit steht in eckigen Klammern. Die
Spalten werden über den Kopf zugeordnet (nicht über feste Positionen) und die Einheit wird gegen
die erwartete Einheit geprüft und bei Abweichung umgerechnet (z. B. mg statt µg).
"""

import re
import time
import zipfile
from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.importers.types import IngredientRecord, Progress

KJ_PER_KCAL = 4.184
SALT_G_PER_G_SODIUM = 2.5  # Salz (NaCl) = Natrium × 2,5

# Fortschritt: alle so viele Zeilen eine Meldung; Schätzung, falls die Datei keine Größe nennt
PROGRESS_EVERY_ROWS = 250
EXPECTED_ROWS = 7140

# ---------------------------------------------------------------------------
# Warengruppen (erster Buchstabe des BLS-Codes)
# ---------------------------------------------------------------------------
# Die Dokumentation (BLS_4_0_Dokumentation_DE.pdf) nennt die Gruppen nicht einzeln; die Zuordnung
# stammt aus den Daten (Bezeichnungen je Buchstabe und Folgeziffern der BLS 4.0 geprüft).
BLS_GROUPS: dict[str, str] = {
    "B": "Brot und Kleingebäck",
    "C": "Getreide und Getreideprodukte",
    "D": "Backwaren, Kuchen und Torten",
    "E": "Eier und Teigwaren",
    "F": "Obst",
    "G": "Gemüse",
    "H": "Hülsenfrüchte, Nüsse, Samen und pflanzliche Alternativen",
    "K": "Kartoffeln, Pilze und Stärke",
    "M": "Milch und Milchprodukte",
    "N": "Alkoholfreie Getränke",
    "P": "Alkoholische Getränke",
    "Q": "Fette und Öle",
    "R": "Gewürze, Würzmittel, Soßen und Zusatzstoffe",
    "S": "Zucker, Süßwaren und Speiseeis",
    "T": "Fisch und Meeresfrüchte",
    "U": "Fleisch",
    "V": "Geflügel, Wild und Innereien",
    "W": "Wurstwaren und Fleischerzeugnisse",
    "X": "Gerichte und Zubereitungen (Suppen, Soßen, Beilagen, Salate)",
    "Y": "Gerichte und Zubereitungen (Fleisch-, Fisch-, Eier- und Süßspeisen)",
}  # fmt: skip

# Gerichte/Zubereitungen: gemischte, fertig zubereitete Einträge, die beim Zuordnen von Zutaten
# schlechter eingestuft werden sollen (z. B. "Rührei" statt der Zutat "Hühnerei roh").
DISH_GROUPS = frozenset("XY")

# Grundprodukt-Gruppen für die zweite Regel der Rohware-Erkennung (siehe `bls_raw_flag`)
_BASIC_GROUPS = frozenset("CEFGHKMNQRSTUV")
_RAW_WORD = re.compile(r"\broh\b", re.IGNORECASE)
_PROCESSED_WORD = re.compile(
    r"gekocht|gebraten|gedünstet|gebacken|geschmort|gegrillt|frittiert|gegart|zubereitet|konserve"
    r"|tiefgefroren|getrocknet|geröstet|gebrannt|gepökelt|geräuchert|gesalzen|gezuckert|gesüßt"
    r"|pulver|instant|dragiert|paniert|gefüllt",
    re.IGNORECASE,
)

# ---------------------------------------------------------------------------
# Spaltenzuordnung
# ---------------------------------------------------------------------------
# BLS-Code -> erwartete Einheit je 100 g. Weicht der Kopf ab, wird umgerechnet.
_UNITS: dict[str, str] = {
    "ENERCC": "kcal", "ENERCJ": "kJ", "PROT625": "g", "FAT": "g", "CHO": "g", "FIBT": "g",
    "NACL": "g", "NA": "mg", "FASAT": "g", "SUGAR": "g", "FAPUN3": "g", "F20:5CN3": "g",
    "F22:6CN3": "g", "K": "mg", "CA": "mg", "MG": "mg", "P": "mg", "FE": "mg", "ZN": "mg",
    "ID": "µg", "VITA": "µg", "VITD": "µg", "VITE": "mg", "VITK": "µg", "VITC": "mg",
    "VITB12": "µg", "FOL": "µg",
}  # fmt: skip

# Mikronährstoffe: unser Schlüssel (`defaults.MICRO_KEYS`) -> BLS-Code. EPA+DHA ist eine Summe
# (`epa_dha_g`) und wird separat gebildet.
MICRO_CODES: dict[str, str] = {
    "sat_fat_g": "FASAT", "sugar_g": "SUGAR", "omega3_g": "FAPUN3", "sodium_mg": "NA",
    "potassium_mg": "K", "calcium_mg": "CA", "magnesium_mg": "MG", "phosphorus_mg": "P",
    "iron_mg": "FE", "zinc_mg": "ZN", "iodine_ug": "ID", "vit_a_ug": "VITA", "vit_d_ug": "VITD",
    "vit_e_mg": "VITE", "vit_k_ug": "VITK", "vit_c_mg": "VITC", "vit_b12_ug": "VITB12",
    "folate_ug": "FOL",
}  # fmt: skip
_EPA_CODE = "F20:5CN3"
_DHA_CODE = "F22:6CN3"

_UNIT_IN_HEADER = re.compile(r"\[\s*([^/\]\s]+)\s*/\s*100\s*g\s*\]")
_MASS_IN_G = {"g": 1.0, "mg": 1e-3, "ug": 1e-6}


def bls_group_name(code: str) -> str | None:
    """Deutscher Name der Warengruppe (erster Buchstabe des BLS-Codes) oder None."""
    code = (code or "").strip()
    return BLS_GROUPS.get(code[:1].upper()) if code else None


def bls_dish_flag(code: str) -> bool:
    """True für Gerichte/Zubereitungen (Gruppen X und Y), die als Zutat schlechter passen."""
    code = (code or "").strip()
    return bool(code) and code[0].upper() in DISH_GROUPS


def bls_raw_flag(code: str, name: str) -> bool:
    """True für Rohware bzw. unverarbeitete Grundprodukte (einfache, bewusst grobe Regel).

    1. Gerichte (X, Y) sind nie Rohware (auch nicht "Hefeteig, roh").
    2. Der Name enthält das Wort "roh" ("Speisezwiebel roh", "Hühnerei roh"); "Rohpökelware",
       "Rohmilch" zählen nicht, weil dort kein eigenständiges Wort "roh" steht.
    3. Grundprodukt: Code endet auf "00" (kein Zubereitungsschritt im Code, z. B. "Olivenöl",
       "Haferflocken", "Vollmilch"), die Gruppe ist eine Zutatengruppe (kein Brot, Gebäck, Wurst,
       Gericht) und der Name nennt keine Verarbeitung (gekocht, getrocknet, Konserve, ...).
    """
    code = (code or "").strip().upper()
    if not code or bls_dish_flag(code):
        return False
    if _RAW_WORD.search(name or ""):
        return True
    return code.endswith("00") and code[0] in _BASIC_GROUPS and not _PROCESSED_WORD.search(name or "")


# ---------------------------------------------------------------------------
# Hilfsfunktionen
# ---------------------------------------------------------------------------


def _norm_unit(unit: str) -> str:
    return unit.strip().lower().replace("µ", "u").replace("μ", "u")


def _unit_factor(code: str, header_unit: str | None) -> float:
    """Faktor, der einen Wert in der Kopf-Einheit in die erwartete Einheit umrechnet."""
    expected = _UNITS[code]
    if header_unit is None:
        return 1.0
    src, dst = _norm_unit(header_unit), _norm_unit(expected)
    if src == dst:
        return 1.0
    if src in _MASS_IN_G and dst in _MASS_IN_G:
        return _MASS_IN_G[src] / _MASS_IN_G[dst]
    raise ValueError(f"BLS-Spalte {code}: Einheit '{header_unit}' statt '{expected}' nicht umrechenbar")


def _num(value: Any) -> float | None:
    """Zahl oder numerischer Text (auch Dezimalkomma); alles andere (leer, "-", Text) → None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        number = float(value)
    elif isinstance(value, str):
        text = value.strip().replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            return None
    else:
        return None
    if number != number or number in (float("inf"), float("-inf")) or number < 0:
        return None  # NaN, unendlich und negative Werte sind keine gültigen Gehalte
    return number


def _parse_header(header: Sequence[Any]) -> dict[str, tuple[int, float]]:
    """Wertspalten des Kopfes: BLS-Code -> (Spaltenindex, Umrechnungsfaktor)."""
    first = str(header[0] or "").lower() if header else ""
    if "code" not in first:
        raise ValueError("Keine BLS-Datendatei: Spalte A heißt nicht 'BLS Code'")
    columns: dict[str, tuple[int, float]] = {}
    for index, cell in enumerate(header):
        if index < 3 or not isinstance(cell, str):
            continue
        text = cell.strip()
        if text.endswith(("Datenherkunft", "Referenz")):
            continue
        code = text.split(" ", 1)[0]
        if code not in _UNITS or code in columns:
            continue
        match = _UNIT_IN_HEADER.search(text)
        columns[code] = (index, _unit_factor(code, match.group(1) if match else None))
    if "ENERCC" not in columns and "ENERCJ" not in columns:
        raise ValueError("Keine BLS-Datendatei: Energiespalten (ENERCC/ENERCJ) fehlen")
    return columns


def _open_xlsx_bytes(path: Path) -> BytesIO:
    """XLSX direkt oder aus der ZIP-Datei (ohne Entpacken auf die Platte) in den Speicher lesen."""
    if not path.is_file():
        raise FileNotFoundError(f"BLS-Datei nicht gefunden: {path}")
    if not zipfile.is_zipfile(path):
        raise ValueError(f"{path.name}: weder XLSX noch ZIP-Datei")
    with zipfile.ZipFile(path) as zf:
        names = zf.namelist()
        if "xl/workbook.xml" in names:  # die Datei selbst ist eine XLSX (ein ZIP-Container)
            return BytesIO(path.read_bytes())
        candidates = [
            info
            for info in zf.infolist()
            if info.filename.lower().endswith(".xlsx")
            and not info.filename.rsplit("/", 1)[-1].startswith(("~$", "."))
            and "__macosx" not in info.filename.lower()
        ]
        if not candidates:
            raise ValueError(f"{path.name}: keine XLSX-Datei in der ZIP-Datei")
        # Die Datendatei heißt "...Daten..."; sonst die größte (nicht die Komponentenliste)
        data = [c for c in candidates if "daten" in c.filename.lower()]
        pool = data or [c for c in candidates if "component" not in c.filename.lower()] or candidates
        chosen = max(pool, key=lambda info: info.file_size)
        return BytesIO(zf.read(chosen))


def _record(row: Sequence[Any], columns: dict[str, tuple[int, float]]) -> IngredientRecord | None:
    """Eine Datenzeile als Zutat; None, wenn Name/Code fehlen oder keine Energie-/Makrowerte da sind."""

    def get(code: str) -> float | None:
        spec = columns.get(code)
        if spec is None or spec[0] >= len(row):
            return None
        value = _num(row[spec[0]])
        return None if value is None else value * spec[1]

    code = str(row[0]).strip() if row and row[0] is not None else ""
    name = " ".join(str(row[1]).split()) if len(row) > 1 and row[1] is not None else ""
    if not code or not name:
        return None

    kcal = get("ENERCC")
    if kcal is None:
        kj = get("ENERCJ")
        kcal = None if kj is None else kj / KJ_PER_KCAL
    protein, fat, carb = get("PROT625"), get("FAT"), get("CHO")
    if kcal is None and protein is None and fat is None and carb is None:
        return None

    sodium_mg = get("NA")
    salt = get("NACL")
    if salt is None and sodium_mg is not None:
        salt = sodium_mg * SALT_G_PER_G_SODIUM / 1000.0

    micros: dict[str, float] = {}
    for key, bls_code in MICRO_CODES.items():
        value = get(bls_code)
        if value is not None:
            micros[key] = value
    epa, dha = get(_EPA_CODE), get(_DHA_CODE)
    if epa is not None or dha is not None:
        micros["epa_dha_g"] = (epa or 0.0) + (dha or 0.0)

    return IngredientRecord(
        name=name,
        source="bls",
        source_code=code,
        category=bls_group_name(code),
        kcal_100=kcal,
        protein_100=protein,
        fat_100=fat,
        carb_100=carb,
        fiber_100=get("FIBT"),
        salt_100=salt,
        micros=micros,
        is_fish=code[0].upper() == "T",
    )


# ---------------------------------------------------------------------------
# Öffentliche Funktionen
# ---------------------------------------------------------------------------


def import_bls_with_stats(
    path: Path, progress: Progress | None = None
) -> tuple[list[IngredientRecord], dict[str, int]]:
    """BLS-Datei (XLSX oder ZIP) lesen. Liefert die Zutaten und die Zählung rows/imported/skipped.

    `rows` = Zeilen mit Inhalt (ohne Kopfzeile), `skipped` = Zeilen ohne Namen/Code oder ohne
    Energie und ohne Protein/Fett/Kohlenhydrate. Fortschritt: `progress(Anteil 0..1, Meldung)`.
    """
    path = Path(path)
    started = time.monotonic()

    def report(fraction: float, message: str) -> None:
        if progress is not None:
            progress(max(0.0, min(1.0, fraction)), message)

    report(0.0, f"BLS: {path.name} wird geöffnet")
    buffer = _open_xlsx_bytes(path)
    workbook = load_workbook(buffer, read_only=True, data_only=True)
    records: list[IngredientRecord] = []
    stats = {"rows": 0, "imported": 0, "skipped": 0}
    try:
        sheet = workbook.worksheets[0]
        rows_iter = sheet.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if header is None:
            raise ValueError("BLS-Datei ist leer")
        columns = _parse_header(header)

        total = sheet.max_row - 1 if isinstance(sheet.max_row, int) and sheet.max_row > 1 else EXPECTED_ROWS
        for row in rows_iter:
            if not row or all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in row):
                continue  # leere Zeile (Dateiende)
            stats["rows"] += 1
            record = _record(row, columns)
            if record is None:
                stats["skipped"] += 1
            else:
                records.append(record)
                stats["imported"] += 1
            if stats["rows"] % PROGRESS_EVERY_ROWS == 0:
                report(
                    min(0.99, stats["rows"] / max(total, stats["rows"] + 1)),
                    f"BLS: {stats['rows']} Zeilen gelesen",
                )
    finally:
        workbook.close()

    seconds = time.monotonic() - started
    report(
        1.0,
        f"BLS: {stats['imported']} Zutaten gelesen, {stats['skipped']} Zeilen übersprungen "
        f"({seconds:.1f} s)",
    )
    return records, stats


def import_bls(path: Path, progress: Progress | None = None) -> list[IngredientRecord]:
    """BLS-Datei (XLSX oder ZIP) lesen und die Zutaten liefern (Zählung: `import_bls_with_stats`)."""
    return import_bls_with_stats(path, progress)[0]
